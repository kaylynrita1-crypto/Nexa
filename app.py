import streamlit as st
from pathlib import Path
css = Path("style.css").read_text(encoding="utf-8")
st.markdown(f"<style>{css}</style>", unsafe_allow_html=True)
st.markdown(
    '<div class="nexa-signature">NEXA by Rita</div>',
    unsafe_allow_html=True
)
import os
from datetime import datetime
from docx import Document
from io import BytesIO
from PIL import Image, ImageDraw, ImageFont
import textwrap                                              
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from dotenv import load_dotenv
from google import genai

load_dotenv()
api_key = os.getenv("GEMINI_API_KEY") or st.secrets.get("GEMINI_API_KEY", None)

client = genai.Client(api_key=api_key) if api_key else None
if "tool_recommendation" not in st.session_state:
    st.session_state["tool_recommendation"] = ""

st.set_page_config(
    page_title="NEXA | Intelligent Digital Toolbox",
    page_icon="💜",
    layout="wide"
)


# ==============================
# NEXA SIDEBAR
# ==============================

with st.sidebar:
    st.markdown("# 💚 NEXA")
    st.caption("Your ideas. Unlocked.")

    st.divider()

    if st.button("＋ New chat", use_container_width=True):
        st.session_state["nexa_prompt"] = ""
        st.rerun()

    st.markdown("### Your tools")

    selected_tool = st.radio(
        "Choose a tool",
        [
            "AI Assistant",
            "CV Creator",
            "Budget Planner",
            "Project Planner",
            "Poster Creator",
        ],
        label_visibility="collapsed",
    )

    st.divider()
    st.caption("NEXA by Rita")
# ==============================
# NEXA SIDEBAR
# ==============================

with st.sidebar:
    st.markdown("# 💚 NEXA")
    st.caption("Your ideas. Unlocked.")

    st.divider()

    if st.button("＋ New chat", use_container_width=True, key="nexa_sidebar_new_chat"):
        st.session_state["nexa_prompt"] = ""
        st.rerun()

    st.markdown("### Your tools")

    selected_tool = st.radio(
    "Choose a tool",
    [
        "AI Assistant",
        "CV Creator",
        "Budget Planner",
        "Project Planner",
        "Poster Creator",
    ],
    label_visibility="collapsed",
    key="nexa_sidebar_tool_selector",
)

    st.divider()

    st.caption("NEXA by Rita")

st.title("💜 NEXA")
st.subheader("Your Intelligent Digital Toolbox")
st.write("Your ideas. Unlocked. Let's turn your idea into something real.")

st.divider()
st.title("💎 NEXA")
st.caption("Your Intelligent Digital Toolbox")
st.header("✨ What would you like to create today?")
st.write(
    "Bring your ideas to life. NEXA helps you create, "
    "plan, organize, and turn your ideas into finished work."
    )
st.markdown("### 💡 Tell NEXA what's on your mind")
idea = st.text_area(
    "Describe your idea",
    placeholder="For example: Help me create a professional CV..."
)

st.write("")
st.caption("✨ NEXA will find the right tools to bring your idea to life.")
if st.button("✨ Unlock My Idea", use_container_width=True):
    if not idea.strip():
        st.warning("Please describe your idea first.")
    elif client is None:
        st.error("NEXA cannot find its API key. Check your .env file.")
    else:
        with st.spinner("NEXA is finding the best tool for you..."):
            try:
                response = client.models.generate_content(
                    model="gemini-3.8-flash",
                    contents=(
                        "You are NEXA, an intelligent digital toolbox. "
                        "Choose the best tool for this user request: "
                        f"{idea}\n\n"
                        
"Choose the best tool or combination of tools for "
"the user's request. Available tools are: CV Creator, "
"Budget Planner, Project Planner, and AI Assistant. "
"Choose CV Creator for CVs and resumes, Budget Planner "
"for expenses and saving money, Project Planner for "
"tasks and deadlines, and AI Assistant for general help. "
"If the user needs multiple tools, list each relevant tool. "
"Briefly explain why the tools fit the request. "
"Start your reply with the exact tool name or names, "
"using these names exactly as written."

                    )
                )

                recommendation = response.text.strip()
                allowed_tools = [
                    "CV Creator",
                    "Budget Planner",
                    "Project Planner",
                    "AI Assistant"
                ]

                if not any(tool in recommendation for tool in allowed_tools):
                    recommendation = "AI Assistant"

                st.session_state["tool_recommendation"] = recommendation
                st.session_state["user_idea"] = idea

            
            except Exception as e:
                idea_lower = idea.lower()

                if "cv" in idea_lower or "resume" in idea_lower:
                    recommendation = "CV Creator"
                elif any(word in idea_lower for word in [
                    "budget", "expenses", "saving", "save money"
                ]):
                    recommendation = "Budget Planner"
                elif any(word in idea_lower for word in [
                    "project", "tasks", "deadline", "organize"
                ]):
                    recommendation = "Project Planner"
                else:
                    recommendation = "AI Assistant"

                st.session_state["tool_recommendation"] = recommendation
                st.session_state["user_idea"] = idea
                st.info("NEXA used its backup recommendation because its AI is busy.")


if "tool_recommendation" in st.session_state:
    st.divider()
    st.subheader("✨ Your NEXA Recommendation")
    st.success(
        "The best tool for your idea is: "
        + st.session_state["tool_recommendation"]
    )
    

   
    recommendation = st.session_state["tool_recommendation"]
    user_idea = st.session_state.get("user_idea", "")

    if "action_plan" not in st.session_state:
        st.session_state["action_plan"] = ""

    if st.button("✨ Create My Personalized Action Plan"):
        if client is None:
            st.warning("NEXA AI is unavailable right now. Please try again later.")
        else:
            with st.spinner("NEXA is creating a plan for your idea..."):
                try:
                    plan_response = client.models.generate_content(
                        model="gemini-3.8-flash",
                        contents=(
                            "You are NEXA, an intelligent digital toolbox. "
                            "Create a practical, personalized action plan for this idea: "
                            + user_idea
                            + "\n\nRecommended tools: "
                            + recommendation
                            + "\n\nWrite 4 to 7 clear, numbered steps. "
                            "Make every step specific to the user's idea. "
                            "Mention which recommended NEXA tool to use for each relevant step. "
                            "Keep the advice realistic and beginner-friendly. "
                            "Do not claim the user has completed any task."
                        )
                    )

                    st.session_state["action_plan"] = plan_response.text
                except Exception as e:
                    st.warning(
                        "Gemini is unavailable. NEXA is creating a backup plan."
                    )

                    idea_text = user_idea.lower()
                    backup_steps = [
                        f"1. Define your goal: {user_idea}",
                        "2. Break your goal into small, manageable tasks.",
                        "3. Identify the resources, time, and money you need.",
                        "4. Use the recommended NEXA tools to organize your work.",
                        "5. Review your progress and adjust your next steps."
                    ]

                    if "Budget Planner" in recommendation:
                        backup_steps.append(
                            "6. Use Budget Planner to estimate costs and track spending."
                        )

                    if "Project Planner" in recommendation:
                        backup_steps.append(
                            "7. Use Project Planner to organize tasks and deadlines."
                        )

                    if "CV Creator" in recommendation:
                        backup_steps.append(
                            "6. Use CV Creator to prepare your professional CV."
                        )

                    st.session_state["action_plan"] = "\n\n".join(backup_steps)
                    st.success("Your backup action plan is ready!")

    if st.session_state["action_plan"]:
        st.markdown(st.session_state["action_plan"])
        st.download_button(
            label="📥 Download My Action Plan",
            data=st.session_state["action_plan"],
            file_name="NEXA_Action_Plan.txt",
            mime="text/plain"
        )

        plan_title = st.text_input(
            "Name your action plan before saving",
            value="My Action Plan",
            key="new_plan_title"
        )

        if st.button("💾 Save My Action Plan"):
            if not plan_title.strip():
                st.warning("Please enter a name for your plan.")
            else:
                saved_date = datetime.now().strftime("%Y-%m-%d %H:%M")

                with open("NEXA_Saved_Plans.txt", "a", encoding="utf-8") as file:
                    file.write("\n\n")
                    file.write(f"TITLE: {plan_title.strip()}\n")
                    file.write(f"SAVED: {saved_date}\n")
                    file.write(st.session_state["action_plan"])
                    file.write("\n" + "=" * 40 + "\n")

                st.success("Your named action plan has been saved! 💜")
                st.divider()
st.subheader("🧰 Your Workspace")
selected_tools = st.multiselect(
    "Your recommended tool is preselected. You can also choose additional tools:",
        [
        "CV Creator",
        "Budget Planner",
        "Project Planner",
        "Poster Creator",
        "AI Assistant"
    ],
    default=[
        tool for tool in [
            "CV Creator",
            "Budget Planner",
            "Project Planner",
            "AI Assistant"
        ]
        if tool in st.session_state["tool_recommendation"]
    ]
)

   
    
if st.button("Open My Chosen Tools"):
    if not selected_tools:
        st.warning("Please select at least one tool.")
    else:
        st.session_state["chosen_tools"] = selected_tools

        if "CV Creator" in selected_tools:
            st.session_state["show_cv"] = True
        else:
            st.session_state["show_cv"] = False

        if "Poster Creator" in selected_tools:
            st.session_state["show_poster"] = True
        else:
            st.session_state["show_poster"] = False

        st.rerun()




if st.session_state.get("show_cv", False):
    st.divider()
    st.header("📄 NEXA CV Creator")
    st.caption("✨ Build a professional CV that's ready to download and share.")

    with st.form("cv_form"):
        full_name = st.text_input("Full name")
        email = st.text_input("Email address")
        phone = st.text_input("Phone number")
        location = st.text_input("Town or city")
        career = st.text_input("Job title or career objective")
        education = st.text_area("Education")
        experience = st.text_area("Work experience")
        skills = st.text_area("Skills")
        submit_cv = st.form_submit_button("Create My CV")

    if submit_cv:
        if not full_name.strip():
            st.error("Please enter your full name.")
        else:
            
            doc = Document()
            # NEXA Emerald & Champagne CV styling
            emerald = RGBColor(18, 92, 74)
            champagne = RGBColor(184, 150, 90)

            # Set the default font
            normal_style = doc.styles["Normal"]
            normal_style.font.name = "Aptos"
            normal_style.font.size = Pt(10)
            normal_style.font.color.rgb = RGBColor(55, 65, 60)

            # Style section headings
            for style_name in ["Heading 1", "Heading 2"]:
                heading_style = doc.styles[style_name]
                heading_style.font.name = "Aptos Display"
                heading_style.font.color.rgb = emerald
                heading_style.font.bold = True

            doc.styles["Heading 1"].font.size = Pt(14)
            doc.styles["Heading 2"].font.size = Pt(12)

            # Give paragraphs comfortable spacing
            normal_style.paragraph_format.space_after = Pt(6)
            # Document margins
            section = doc.sections[0]
            section.top_margin = 720000
            section.bottom_margin = 720000
            section.left_margin = 850000
            section.right_margin = 850000

            # 
            # NEXA professional name and contact header
            title = doc.add_heading(full_name.strip(), level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            title_run = title.runs[0]
            title_run.font.name = "Aptos Display"
            title_run.font.size = Pt(26)
            title_run.font.bold = True
            title_run.font.color.rgb = emerald

            # Elegant champagne divider
            divider = doc.add_paragraph()
            divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
            divider.paragraph_format.space_after = Pt(8)

            divider_run = divider.add_run(
                "━━━━━━━━━━━━━━━━━━━━━━━━"
            )
            divider_run.font.color.rgb = champagne
            divider_run.font.size = Pt(10)

            # Contact details
            contact = doc.add_paragraph()
            contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact.paragraph_format.space_after = Pt(18)

            contact_run = contact.add_run(
                "  |  ".join(
                    item.strip()
                    for item in [email, phone, location]
                    if item.strip()
                )
            )
            contact_run.font.name = "Aptos"
            contact_run.font.size = Pt(9)
            contact_run.font.color.rgb = RGBColor(85, 95, 90)

            # Career objective
            doc.add_heading("Professional Profile", level=1)
            doc.add_paragraph(career.strip() or "Not provided")

            # Education
            doc.add_heading("Education", level=1)
            doc.add_paragraph(education.strip() or "Not provided")

            # Work experience
            doc.add_heading("Work Experience", level=1)
            doc.add_paragraph(experience.strip() or "Not provided")

            # Skills
            doc.add_heading("Skills", level=1)
            doc.add_paragraph(skills.strip() or "Not provided")

            # Prepare Word file for download
            word_file = BytesIO()
            doc.save(word_file)
            word_file.seek(0)

            st.success("Your CV has been created!")

            st.download_button(
                "⬇️ Download Your CV (Word)",
                data=word_file.getvalue(),
                file_name="NEXA_CV.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

            # Show a simple preview in NEXA
            st.subheader("📄 CV Preview")
            st.markdown(f"## {full_name.strip()}")
            st.caption(f"{email} · {phone} · {location}")
            st.markdown("### Professional Profile")
            st.write(career or "Not provided")
            st.markdown("### Education")
            st.write(education or "Not provided")
            st.markdown("### Work Experience")
            st.write(experience or "Not provided")
            st.markdown("### Skills")
            st.write(skills or "Not provided")


if st.session_state.get("show_poster", False):
    st.divider()
    st.header("🎨 NEXA Poster Creator")
    st.caption("✨ Turn your idea into a beautiful, downloadable poster.")

    with st.form("poster_form"):
        poster_title = st.text_input("Poster title")
        poster_subtitle = st.text_input("Subtitle or slogan")
        poster_details = st.text_area("Extra details")
        poster_style = st.selectbox(
            "Choose your poster style",
            ["Emerald & Champagne", "Dark & Elegant", "Bright & Vibrant"]
        )
        poster_submit = st.form_submit_button("Create My Poster")


    
    
    if poster_submit:
        if not poster_title.strip():
            st.error("Please enter a poster title.")
        else:
            width, height = 1080, 1350

            palettes = {
                "Emerald & Champagne": ((18, 92, 74), (232, 215, 174), (255, 255, 255)),
                "Dark & Elegant": ((35, 25, 45), (218, 190, 145), (255, 255, 255)),
                "Bright & Vibrant": ((35, 90, 180), (255, 210, 60), (255, 255, 255)),
            }

            background, accent, foreground = palettes[poster_style]
            poster = Image.new("RGB", (width, height), background)
            draw = ImageDraw.Draw(poster)

            # Decorative frame
            draw.rounded_rectangle(
                (45, 45, width - 45, height - 45),
                radius=35,
                outline=accent,
                width=6
            )

            # Load fonts
            try:
                title_font = ImageFont.truetype("arialbd.ttf", 76)
                subtitle_font = ImageFont.truetype("arial.ttf", 38)
                details_font = ImageFont.truetype("arial.ttf", 30)
                brand_font = ImageFont.truetype("arialbd.ttf", 28)
            except OSError:
                title_font = ImageFont.load_default()
                subtitle_font = ImageFont.load_default()
                details_font = ImageFont.load_default()
                brand_font = ImageFont.load_default()

            def draw_centered_wrapped(text, y, font, fill, max_chars):
                lines = textwrap.wrap(text, width=max_chars) or [""]
                for line in lines:
                    box = draw.textbbox((0, 0), line, font=font)
                    line_width = box[2] - box[0]
                    draw.text(
                        ((width - line_width) / 2, y),
                        line,
                        font=font,
                        fill=fill
                    )
                    y += (box[3] - box[1]) + 18
                return y

            y = 250
            y = draw_centered_wrapped(
                poster_title.strip(), y, title_font, accent, 20
            )

            if poster_subtitle.strip():
                y += 35
                y = draw_centered_wrapped(
                    poster_subtitle.strip(), y, subtitle_font, foreground, 34
                )

            if poster_details.strip():
                y += 60
                y = draw_centered_wrapped(
                    poster_details.strip(), y, details_font, foreground, 45
                )

            draw.line((200, height - 190, width - 200, height - 190), fill=accent, width=3)
            draw.text((width // 2, height - 145), "CREATED WITH NEXA", font=brand_font, fill=accent, anchor="mm")

            image_buffer = BytesIO()
            poster.save(image_buffer, format="PNG")
            poster_bytes = image_buffer.getvalue()

            st.success("Your poster is ready!")
            st.image(poster_bytes, caption="Your NEXA poster", use_container_width=True)

            st.download_button(
                "⬇️ Download My Poster (PNG)",
                data=poster_bytes,
                file_name="NEXA_Poster.png",
                mime="image/png",
                use_container_width=True
            )


# ==============================
# NEXA BUDGET PLANNER
# ==============================
if "Budget Planner" in st.session_state.get("chosen_tools", []):
    st.divider()
    st.header("💰 NEXA Budget Planner")
    st.caption("✨ Take control of your money with a clearer plan for spending and saving.")
    with st.form("budget_form"):
        st.subheader("Your monthly finances")

        currency = st.selectbox(
            "Currency",
            ["KES (KSh)", "USD ($)", "GBP (£)", "EUR (€)"]
        )

        income = st.number_input(
            "Monthly income",
            min_value=0.0,
            value=0.0,
            step=500.0
        )

        st.subheader("Monthly expenses")

        housing = st.number_input("Rent / Housing", min_value=0.0, value=0.0, step=100.0)
        food = st.number_input("Food / Shopping", min_value=0.0, value=0.0, step=100.0)
        transport = st.number_input("Transport", min_value=0.0, value=0.0, step=100.0)
        utilities = st.number_input("Bills / Utilities", min_value=0.0, value=0.0, step=100.0)
        other = st.number_input("Other expenses", min_value=0.0, value=0.0, step=100.0)

        submitted_budget = st.form_submit_button("Calculate My Budget")

    if submitted_budget:
        expenses = housing + food + transport + utilities + other
        balance = income - expenses

        currency_symbol = {
            "KES (KSh)": "KSh",
            "USD ($)": "$",
            "GBP (£)": "£",
            "EUR (€)": "€"
        }[currency]

        st.divider()
        st.subheader("📊 Your monthly budget results")

        col1, col2, col3 = st.columns(3)

        col1.metric(
            "Monthly income",
            f"{currency_symbol} {income:,.2f}"
        )

        col2.metric(
            "Total expenses",
            f"{currency_symbol} {expenses:,.2f}"
        )

        col3.metric(
            "Money remaining",
            f"{currency_symbol} {balance:,.2f}"
        )

        if balance > 0:
            st.success(
                f"You have {currency_symbol} {balance:,.2f} remaining this month."
            )
            st.write(
                f"Suggested savings target (20% of income): "
                f"{currency_symbol} {income * 0.20:,.2f}"
            )
        elif balance == 0:
            st.warning("Your income is fully allocated. There is no money left over.")
        else:
            st.error(
                f"You are over budget by {currency_symbol} {abs(balance):,.2f}. "
                "Consider reviewing your expenses."
            )

        st.subheader("Expense breakdown")
        st.write(f"Housing: {currency_symbol} {housing:,.2f}")
        st.write(f"Food: {currency_symbol} {food:,.2f}")
        st.write(f"Transport: {currency_symbol} {transport:,.2f}")
        st.write(f"Bills: {currency_symbol} {utilities:,.2f}")
        st.write(f"Other: {currency_symbol} {other:,.2f}")



# ==============================
# NEXA PROJECT PLANNER
# ==============================

if "Project Planner" in st.session_state.get("chosen_tools", []):
    st.divider()
    st.header("📋 NEXA Project Planner")
    st.caption("✨ Turn big ideas into manageable tasks, clear priorities, and an actionable plan.")

    with st.form("project_planner_form"):
        project_name = st.text_input(
            "What is your project called?",
            placeholder="e.g. Start an online clothing business"
        )

        project_goal = st.text_area(
            "What do you want to achieve?",
            placeholder="Describe what success looks like..."
        )

        deadline = st.date_input(
            "Target completion date"
        )

        st.subheader("Plan your tasks")

        task1 = st.text_input("Task 1", placeholder="e.g. Research my target customers")
        priority1 = st.selectbox("Task 1 priority", ["High", "Medium", "Low"])

        task2 = st.text_input("Task 2", placeholder="e.g. Calculate startup costs")
        priority2 = st.selectbox("Task 2 priority", ["High", "Medium", "Low"])

        task3 = st.text_input("Task 3", placeholder="e.g. Choose suppliers")
        priority3 = st.selectbox("Task 3 priority", ["High", "Medium", "Low"])

        task4 = st.text_input("Task 4", placeholder="e.g. Prepare a marketing plan")
        priority4 = st.selectbox("Task 4 priority", ["High", "Medium", "Low"])

        task5 = st.text_input("Task 5", placeholder="e.g. Launch the project")
        priority5 = st.selectbox("Task 5 priority", ["High", "Medium", "Low"])

        create_project = st.form_submit_button("✨ Create My Project Plan")

    if create_project:
        if not project_name.strip():
            st.error("Please enter a project name.")
        else:
            tasks = [
                (task1, priority1),
                (task2, priority2),
                (task3, priority3),
                (task4, priority4),
                (task5, priority5),
            ]

            project_plan = f"""# {project_name}

## Project Goal
{project_goal}

## Target Completion Date
{deadline}

## Project Tasks
"""

            task_count = 0

            for index, (task, priority) in enumerate(tasks, start=1):
                if task.strip():
                    project_plan += (
                        f"{index}. [ ] {task.strip()} "
                        f"(Priority: {priority})\n"
                    )
                    task_count += 1

            if task_count == 0:
                st.warning("Please enter at least one task.")
            else:
                st.success("Your project plan is ready!")
                st.write(f"**Project:** {project_name}")
                st.write(f"**Tasks included:** {task_count}")
                st.write(f"**Deadline:** {deadline}")

                st.download_button(
                    "⬇️ Download My Project Plan",
                    data=project_plan,
                    file_name="nexa_project_plan.md",
                    mime="text/markdown"
                )

                with st.expander("View your project plan"):
                    st.markdown(project_plan)


if "AI Assistant" in st.session_state.get("chosen_tools", []):
    st.divider()
    st.header("✨ NEXA AI Assistant")
st.write("Describe what you need, and NEXA will help you work it out.")

ai_prompt = st.text_area(
    "What would you like NEXA to help with?",
    key="nexa_ai_prompt",
    placeholder="Example: Help me plan a small business with a limited budget."
)

if st.button("Ask NEXA AI"):
    if not ai_prompt.strip():
        st.warning("Please describe what you need.")
    elif client is None:
        st.error("NEXA cannot find its API key. Check your .env file.")
    else:
        with st.spinner("NEXA is thinking..."):
            try:
                response = client.models.generate_content(
                    model="gemini-3.8-flash",
                    contents=ai_prompt
                )
                st.markdown(response.text)
            except Exception as e:
                st.error(f"AI request failed: {e}")

                # ==============================
# NEXA SAVED PLANS LIBRARY
# ==============================

st.divider()
st.header("📚 My Saved Plans")

if os.path.exists("NEXA_Saved_Plans.txt"):
    with open("NEXA_Saved_Plans.txt", "r", encoding="utf-8") as file:
        saved_plans_text = file.read()

    saved_plans = [
        plan.strip()
        for plan in saved_plans_text.split("=" * 40)
        if plan.strip()
    ]

    if saved_plans:
        st.write("Your previously saved action plans:")

        for number, plan in enumerate(reversed(saved_plans), start=1):
            lines = plan.splitlines()
            title = next(
                (line.replace("TITLE:", "", 1).strip() for line in lines if line.startswith("TITLE:")),
                f"Saved Plan {number}",
            )
            saved_date = next(
                (line.replace("SAVED:", "", 1).strip() for line in lines if line.startswith("SAVED:")),
                "Date not recorded",
            )

            with st.expander(f"📄 {title} — {saved_date}"):
                st.text(plan)
                if st.button("🔄 Reuse This Plan", key=f"reuse_plan_{number}"):
                    st.session_state["action_plan"] = plan
                    st.session_state["reused_plan"] = True
                    st.success("Your saved plan is ready to reuse!")
                    st.rerun()

                st.download_button(
                    label="📥 Download This Plan",
                    data=plan,
                    file_name=f"NEXA_Saved_Plan_{number}.txt",
                    mime="text/plain",
                    key=f"saved_plan_download_{number}"
                )
    else:
        st.info("Your saved plans will appear here.")
else:
    st.info("You haven't saved any plans yet. Generate a plan and click Save My Action Plan!")


# Show a reused plan clearly
if st.session_state.get("reused_plan", False):
    st.divider()
    st.subheader("🔄 Your Reused Action Plan")
    st.markdown(st.session_state["action_plan"])

    st.download_button(
        label="📥 Download Reused Plan",
        data=st.session_state["action_plan"],
        file_name="NEXA_Reused_Action_Plan.txt",
        mime="text/plain",
        key="download_reused_plan"
    )